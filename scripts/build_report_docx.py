"""Create the CA3 Word report from the implemented scope and measured evidence."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
r = json.loads((ROOT/'reports/validation_v3.json').read_text(encoding='utf-8'))
b = json.loads((ROOT/'reports/browser_smoke.json').read_text(encoding='utf-8'))
d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.7)
sec.bottom_margin = Inches(1.0)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
    style = d.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(7)
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.line_spacing = 1.12
d.styles['Title'].font.size = Pt(25)
d.styles['Title'].font.bold = True
d.styles['Subtitle'].font.size = Pt(12)
for name, size in [('Heading 1', 16), ('Heading 2', 12)]:
    d.styles[name].font.size = Pt(size)
    d.styles[name].font.bold = True
    d.styles[name].paragraph_format.space_before = Pt(12)
    d.styles[name].paragraph_format.keep_with_next = True
d.styles['Caption'].font.size = Pt(9)
d.styles['Caption'].font.italic = True
d.core_properties.title = 'ProcureGuard AI CA3 Project Report'
d.core_properties.author = 'Rishi Dwivedi and Siddarth Singh Kushwah'
d.core_properties.subject = 'Agentic invoice exception resolution proof of concept'
header = sec.header.paragraphs[0]
header.text = 'ProcureGuard AI  |  CA3 Project Report'
header.style = d.styles['Normal']
header.runs[0].font.size = Pt(9)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('AIML C  |  SIT Pune  |  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(9)

def p(text): return d.add_paragraph(text)
def h(text, level=1): return d.add_heading(text, level)
def page(title):
    heading=h(title)
    heading.paragraph_format.page_break_before=True
def table(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for cell, width in zip(t.columns, widths): cell.width=Inches(width)
    for cell, text in zip(t.rows[0].cells, headers): cell.text=text
    for row in rows:
        cells=t.add_row().cells
        for cell, text in zip(cells,row): cell.text=str(text)
    for i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if i==0: pr.append(OxmlElement('w:tblHeader'))
        pr.append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'172A42' if i==0 else ('F2F5F8' if i%2==0 else 'FFFFFF'));tc.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tc.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');margins.append(el)
            tc.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1.05
                for run in para.runs:
                    run.font.size=Pt(10)
                    if i==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t

brand=d.add_paragraph()
brand.alignment=WD_ALIGN_PARAGRAPH.CENTER
brand.paragraph_format.space_after=Pt(4)
logo=brand.add_run().add_picture(str(ROOT/'docs/assets/sit_pune_official_logo.png'),width=Inches(6.0))
logo._inline.docPr.set('descr','Symbiosis Institute of Technology official logo and institution name')
d.add_paragraph('ProcureGuard AI CA3 Project Report', 'Title')
d.add_paragraph('Agentic invoice exception resolution', 'Subtitle')
p('CA3 project submission  |  AIML C  |  19 October 2026')
repository=d.add_paragraph()
repository.paragraph_format.space_after=Pt(8)
label=repository.add_run('GitHub repository  ')
label.bold=True;label.font.size=Pt(10)
link=OxmlElement('w:hyperlink')
link.set(qn('r:id'),repository.part.relate_to('https://github.com/siddartttk/ProcureGuard-AI-CA3',RT.HYPERLINK,is_external=True))
link_run=OxmlElement('w:r');properties=OxmlElement('w:rPr')
color=OxmlElement('w:color');color.set(qn('w:val'),'0563C1');properties.append(color)
underline=OxmlElement('w:u');underline.set(qn('w:val'),'single');properties.append(underline)
size=OxmlElement('w:sz');size.set(qn('w:val'),'20');properties.append(size)
link_run.append(properties)
link_text=OxmlElement('w:t');link_text.text='https://github.com/siddartttk/ProcureGuard-AI-CA3'
link_run.append(link_text);link.append(link_run);repository._p.append(link)
table(['Student','PRN'], [('Rishi Dwivedi','24070126152'),('Siddarth Singh Kushwah','24070126174')], [4.3,2.6])
team_spacer=d.paragraphs[-1]
team_spacer.paragraph_format.space_after=Pt(0)
team_spacer.paragraph_format.line_spacing=Pt(2)
team_spacer.add_run().font.size=Pt(2)
h('Project summary')
p('ProcureGuard AI is a working Python and Streamlit proof of concept for investigating invoice exceptions. It selects relevant evidence, validates the source, repeats accounting checks and saves a traceable recommendation for human review. The application remains usable without a language model or network after its dependencies are installed.')
p('The project addresses the work that follows an invoice mismatch. An accounts payable reviewer must establish whether a receipt exists, a price amendment was approved or supplier bank details are authoritative. Those facts belong to different owners. Missing or conflicting records can otherwise lead to repeated handoffs or an unsupported recommendation.')
h('Implemented outcome')
p('The final application supports bounded evidence investigations, an independently implemented accounting critic, SQLite case history, receipt-driven resumption, recovery after a fresh browser session and invoice lifecycle classifications. The workbench exposes original records, verified source evidence and effective investigation snapshots.')
p(f'The recorded validation passed {r["regression_tests"]["passed"]} automated tests and 44 operational checks. The critic-enabled workflow agreed with all 25 fixed development challenge decisions. Four known cases also ran with an installed local model. These synthetic results demonstrate implemented behavior; they do not establish field accuracy, financial savings or production readiness.')
h('Supported scope')
p('Inputs are single-line positive INR records for invoices, purchase orders, goods receipts and vendors. Sources and reviewer roles are fictional. The application produces ready for human review, hold or insufficient data. It contains no payment capability.')

page('Architecture and agent responsibilities')
p('Six logical roles collaborate through an explicit controller. They are not six language models or independently deployed services. Accounting calculations remain deterministic; the optional model selects only a permitted read-only evidence lookup.')
diagram=d.add_picture(str(ROOT/'docs/architecture_overview.png'),width=Inches(6.85))
diagram._inline.docPr.set('descr','Investigation flow showing matching, critic, planner, evidence validation, human review and SQLite case history')
d.add_paragraph('Figure 1  Implemented investigation flow with shared snapshots and SQLite case history', 'Caption')
h('Control flow',2)
p('The controller copies the source dataset and runs matching, risk checks and the independent critic. Verified findings determine the eligible tools. The planner selects one source, and the evidence agent validates and applies accepted records to a copied snapshot. Accounting checks and the critic run again before the next selection. The decision coordinator gates and persists the result.')
h('Bounded model authority',2)
p('A run allows at most three evidence lookups and one accounting recheck. A source tool is attempted at most once per investigation. The model cannot edit invoices, calculate the final decision, exceed the budget or execute payment. Invalid or unavailable model output leads to a visibly labelled deterministic fallback.')

page('Reconciliation and persistent case controls')
h('Numerical policy',2)
p('Amounts use Decimal arithmetic. Illustrative price and arithmetic tolerances are inclusive at INR 0.01. A unit difference that fits the tolerance is still subject to an INR 1.00 line exposure cap. Quantities have no tolerance. Exposure of at least INR 100000 requires human review. Severity weights are illustrative indicators rather than calibrated fraud probabilities.')
h('Cumulative commitments',2)
p('Submitted and recorded invoices sharing the vendor, purchase order and item consume the available order and receipt capacity. Rejected, cancelled, duplicate and superseded submissions are excluded. Recorded means a procurement commitment, not payment approval or settlement. An inactive current submission cannot become ready for review.')
p('Identical active copies with the same vendor invoice number represent one quantity obligation. Conflicting duplicate quantities, prices or totals block confidence in the ledger. Unknown peer lifecycle values also block reconciliation instead of silently releasing capacity. An overcommitted group is held conservatively rather than allocated automatically.')
h('Human lifecycle classification',2)
p('A classification change requires a supported transition, a simulated reviewer role, an expected current status and a rationale. Duplicate and superseded classifications require a related active replacement. Classification changes persist in SQLite and invalidate existing recommendations, including results displayed by another browser session. The remaining obligations must be investigated again.')
h('Recovery and integrity',2)
p('SQLite transactions and expected revisions prevent a stale review from overwriting a newer case. Idempotency keys prevent repeated saves of the same request. Original and effective snapshots, tool observations, findings and review events are retained. Earlier investigation revisions remain downloadable.')
p('Reopening restores a saved packet only when the original source hash matches current data. It preserves the run identifier and restores the current case revision without executing tools again. Changed source data requires revalidation. Snapshot digests and a hash-linked event log detect changes within the trusted local boundary; they are not digital signatures or immutable external storage.')

page('Demonstration and evolving investigations')
h('Two gaps in one case',2)
p('PG103 starts with an invoice for 10 units at INR 120 each, a purchase order at INR 100 and no matching receipt. The offline planner retrieves a procurement-approved amendment first. Price then reconciles, but delivery remains unverified. A warehouse lookup supplies the receipt for 10 units. The critic reruns before the case becomes ready for human review. The original invoice is unchanged.')
p('Refreshing the browser and reopening the workbench restores that saved investigation. The case selector, packet and revision refer to the same case. Reopening performs no new evidence lookups. Verified document values can be inspected in the Evidence provenance tab.')
h('Delayed evidence and resumption',2)
p('PG104 initially lacks an available receipt and remains insufficient data. A reviewer records request_recheck with a rationale. When the fictional delayed warehouse receipt is made available, the same case is investigated again. Mandatory checks and independent verification run on the new effective snapshot. The executed example becomes ready for human review at case revision 3; the before and after packets remain available.')
h('Capacity after a lifecycle decision',2)
p('PG105 and PG105B jointly exceed the same order and receipt capacity, so the group is held. A reviewer can reject PG105B with a rationale. This removes its inactive commitment and invalidates stale recommendations. A fresh investigation of PG105 can then reconcile. PG105B itself remains inactive. Explicit resubmission restores the original classroom demo branch.')
h('Cases that must stop',2)
p('PG106 contains a changed source document that fails its pinned hash check. PG107 contains conflicting authorised prices of INR 120 and INR 130. Both are held instead of accepting an unsupported source. PG110 has no available receipt, so the application displays the missing evidence and the next owner rather than inventing delivery.')
p('The real Chrome interaction check captured 10 application views with no Streamlit exceptions or browser page errors. It exercised saved-case recovery, lifecycle classification, revalidation and source inspection. The timed presentation script targets 4 minutes 45 seconds; this is a written schedule, not a measured spoken rehearsal.')

page('Validation results and limitations')
p('The main operational corpus contains 43 fictional invoices and 44 expected decisions, including the delayed-evidence case before and after availability. The recorded run passed 172 automated tests with zero failures, errors or skips and matched all 44 operational decisions. Tests cover reconciliation, source trust, planner bounds, recovery, stale actions, lifecycle persistence, duplicate allocation and tampering.')
h('Frozen development comparison',2)
p('A separate 25-case fixture and integrity digest were frozen before the v3 changes. Expected labels stay outside the operational Dataset. The same development assistant authored the code and fixtures, so this is development evaluation rather than independently held-out evidence. Two cases deliberately omit a required producer check to test critic containment.')
table(['Evaluation workflow','Correct\ndecisions','Incorrect\nready','Critic\ninterventions'],[( 'Single invoice producer','19 of 25','6','0'),('Historical producer','23 of 25','2','0'),('History with critic','25 of 25','0','4')],[3.1,1.25,1.2,1.35])
p('All three modes produced zero false holds or gaps on these fixtures. Four critic interventions represent two attempts on each of the two injected fault cases. The producer baselines are evaluation-only paths and are not offered as an option that bypasses verification in the application.')
h('Actual local model execution',2)
p('The already installed qwen2.5-coder:7b model ran through Ollama on PG103, PG106, PG107 and PG109. The record contains five actual model selections, zero fallback selections and four expected final decisions. Complete requests, parsed responses and packets are saved. No model was downloaded and no paid model service was used. This small known-case run does not establish that model planning outperforms a rule-based planner.')
h('Limits and future work',2)
p('Supplier and approval records are simulated, and selecting a reviewer role is not authentication. Shared schemas, policies and original data remain common failure modes. A local administrator could rewrite the trusted database boundary. Multi-line invoices, digital PDF or OCR extraction, taxes, credit notes, settlement history and ERP connections remain future work. Independent unseen evaluation and authenticated upstream sources are the next validation priorities.')

page('Running and reviewing the submission')
h('Local setup',2)
p('Use Python 3.11 or newer. On Windows, run setup.ps1 from the project folder to create the environment and install requirements-tested.txt. The first installation requires internet access for packages. Then run START_DEMO.cmd, or run python -m streamlit run app.py from the prepared environment. The default offline mock planner needs no model or network.')
p('Optional model mode requires an already installed compatible Ollama model and a running local service. The demonstrated model is qwen2.5-coder:7b. PROCUREGUARD_OLLAMA_MODEL can select another installed model. The sidebar labels offline and local-model modes separately; a failed model call remains visible as a fallback.')
h('Submission contents',2)
table(['Deliverable','Repository location'], [('Eight slide PowerPoint','presentation/ProcureGuard_AI_CA3_Improved.pptx'),('Word project report','docs/ProcureGuard_AI_CA3_Report.docx'),('Application and engine','app.py and procureguard/'),('Synthetic data and tests','data/ and tests/'),('Measured results and raw evidence','reports/'),('Demo script and viva preparation','presentation/5_minute_script.md and viva_questions.md')],[2.35,4.55])
h('Reproducing checks',2)
p('Run scripts/validate_v3.py with the project Python environment to regenerate test, operational, challenge and evolving-case evidence. scripts/run_live_planner.py separately records actual model execution when the installed service is available. scripts/package_submission.py creates the portable archive and verifies per-file checksums. Runtime databases, environments and private build files are excluded.')
h('Evidence sources',2)
p('This report is based on reports/validation_v3.json, reports/validation_v2.json, reports/pytest-results.xml, reports/live_planner_run.json, reports/evolving_case_demo.json and reports/browser_smoke.json. The architecture and control contracts are documented in docs/architecture.mmd and docs/agents_and_orchestration.md. All reported results describe the recorded implementation on 19 October 2026.')
p('The repository remains private. Faculty need repository access to review the submitted URL.')

output=ROOT/'docs/ProcureGuard_AI_CA3_Report.docx'
d.save(output)
print(output)
